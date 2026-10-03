from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook, load_workbook


def save_result(result: dict, results_dir: str = "../results") -> str:
    """Write the JSON result file for one experiment."""
    exp_id = result["cfg"]["exp_id"]
    out_dir = Path(results_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"{exp_id}.json"
    payload = {"cfg": result["cfg"], "history": result["history"], "summary": result["summary"]}
    with open(path, "w", encoding="utf-8") as f:
        json.dump(payload, f, indent=2, sort_keys=True)
    return str(path)


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    cfg = result["cfg"]
    summary = result["summary"]
    row = {
        "exp_id": cfg["exp_id"],
        "group": cfg["group"],
        "description": cfg["description"],
        "loss": cfg["loss"],
        "optimizer": cfg["optimizer"],
        "lr": cfg["lr"],
        "weight_decay": cfg["weight_decay"],
        "batch": cfg["batch"],
        "epochs": cfg["epochs"],
        "hidden": str(cfg["hidden"]),
        "dropout": cfg["dropout"],
        "clip_norm": cfg["clip_norm"],
        "precision": cfg["precision"],
        "init": cfg["init"],
        "seed": cfg["seed"],
        "step0_loss": summary.get("step0_loss"),
        "best_val_loss": summary.get("best_val_loss"),
        "best_epoch": summary.get("best_epoch"),
        "final_train_loss": summary.get("final_train_loss"),
        "final_val_loss": summary.get("final_val_loss"),
        "val_acc": summary.get("val_acc"),
        "val_macro_f1": summary.get("val_macro_f1"),
        "time_per_epoch_s": summary.get("time_per_epoch_s"),
        "peak_mem_MB": summary.get("peak_mem_MB"),
        "diverged": summary.get("diverged", False),
        "eval_acc": None if eval_scores is None else eval_scores.get("accuracy"),
        "eval_macro_f1": None if eval_scores is None else eval_scores.get("macro_f1"),
        "figure_file": f"figures/{cfg['exp_id']}.png",
        "notes": notes,
    }
    return row


def write_xlsx(rows: list[dict], template_path: str, out_path: str) -> None:
    """Populate the template sheet with the experiment rows."""
    template = Path(template_path)
    if template.exists():
        wb = load_workbook(template)
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Experiments"
    if "Experiments" not in wb.sheetnames:
        ws = wb.create_sheet("Experiments")
    else:
        ws = wb["Experiments"]

    header = [
        "exp_id", "group", "description", "loss", "optimizer", "lr", "weight_decay", "batch", "epochs",
        "hidden", "dropout", "clip_norm", "precision", "init", "seed", "step0_loss", "best_val_loss",
        "best_epoch", "final_train_loss", "final_val_loss", "val_acc", "val_macro_f1", "time_per_epoch_s",
        "peak_mem_MB", "diverged", "eval_acc", "eval_macro_f1", "figure_file", "notes",
    ]
    ws.append(header)
    for row in rows:
        values = [row.get(h) for h in header]
        ws.append(values)
    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
